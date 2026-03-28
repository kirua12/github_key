import re
import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime, timedelta
from collections import Counter

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def parse_chat_log(file_path, days=30):
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()

    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    start_date = today - timedelta(days=days)

    date_pattern = re.compile(
        r"-{2,}\s+(\d{4})년\s+(\d{1,2})월\s+(\d{1,2})일\s+\S+요일\s+-{2,}"
    )
    msg_pattern = re.compile(
        r"^\[([^\]]+)\]\s+\[(?:오전|오후)\s+\d{1,2}:\d{2}\]\s+"
    )

    current_date = None
    user_counter = Counter()

    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue

        date_match = date_pattern.search(line)
        if date_match:
            year = int(date_match.group(1))
            month = int(date_match.group(2))
            day = int(date_match.group(3))
            current_date = datetime(year, month, day)
            continue

        if current_date is None:
            continue
        if current_date < start_date or current_date > today:
            continue

        msg_match = msg_pattern.match(line)
        if msg_match:
            sender_raw = msg_match.group(1)
            if sender_raw in ("오픈채팅봇",):
                continue

            parts = sender_raw.split("/")
            if len(parts) >= 2:
                real_name = parts[0].strip()
                game_nick = parts[1].strip()
                display = f"{real_name}/{game_nick}"
            else:
                display = sender_raw.strip()

            user_counter[display] += 1

    return user_counter


def create_excel(user_counter, output_path, days=30):
    wb = Workbook()
    ws = wb.active
    ws.title = "채팅 통계"

    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    today = datetime.now()
    start = today - timedelta(days=days)
    ws.merge_cells("A1:C1")
    title_cell = ws["A1"]
    title_cell.value = f"채팅 통계 ({start.strftime('%Y-%m-%d')} ~ {today.strftime('%Y-%m-%d')})"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    headers = ["순위", "사용자 (실명/게임닉)", "메시지 수"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    sorted_users = user_counter.most_common()
    for rank, (user, count) in enumerate(sorted_users, 1):
        row = rank + 3
        ws.cell(row=row, column=1, value=rank).border = thin_border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=user).border = thin_border
        ws.cell(row=row, column=3, value=count).border = thin_border
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="center")

    total_row = len(sorted_users) + 4
    ws.cell(row=total_row, column=1, value="").border = thin_border
    total_label = ws.cell(row=total_row, column=2, value="합계")
    total_label.font = Font(bold=True)
    total_label.border = thin_border
    total_val = ws.cell(row=total_row, column=3, value=sum(user_counter.values()))
    total_val.font = Font(bold=True)
    total_val.alignment = Alignment(horizontal="center")
    total_val.border = thin_border

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 12

    wb.save(output_path)


class ChatStatsApp:
    def __init__(self, root):
        self.root = root
        self.root.title("카카오톡 채팅 통계")
        self.root.geometry("700x550")
        self.root.resizable(True, True)

        self.user_counter = None
        self.input_file = None

        self._build_ui()

    def _build_ui(self):
        # --- 상단: 파일 선택 영역 ---
        top_frame = ttk.LabelFrame(self.root, text="파일 선택", padding=10)
        top_frame.pack(fill="x", padx=10, pady=(10, 5))

        self.file_label = ttk.Label(top_frame, text="선택된 파일 없음", foreground="gray")
        self.file_label.pack(side="left", fill="x", expand=True)

        ttk.Button(top_frame, text="txt 파일 열기", command=self._open_file).pack(side="right")

        # --- 중단: 설정 영역 ---
        opt_frame = ttk.LabelFrame(self.root, text="설정", padding=10)
        opt_frame.pack(fill="x", padx=10, pady=5)

        ttk.Label(opt_frame, text="조회 기간 (일):").pack(side="left")
        self.days_var = tk.IntVar(value=30)
        days_spin = ttk.Spinbox(opt_frame, from_=1, to=365, textvariable=self.days_var, width=6)
        days_spin.pack(side="left", padx=(5, 20))

        ttk.Button(opt_frame, text="분석하기", command=self._analyze).pack(side="left")
        ttk.Button(opt_frame, text="엑셀 저장", command=self._save_excel).pack(side="right")

        # --- 하단: 결과 테이블 ---
        table_frame = ttk.LabelFrame(self.root, text="결과 미리보기", padding=10)
        table_frame.pack(fill="both", expand=True, padx=10, pady=(5, 10))

        columns = ("rank", "user", "count")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=15)
        self.tree.heading("rank", text="순위")
        self.tree.heading("user", text="사용자 (실명/게임닉)")
        self.tree.heading("count", text="메시지 수")
        self.tree.column("rank", width=60, anchor="center")
        self.tree.column("user", width=350)
        self.tree.column("count", width=100, anchor="center")

        scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)

        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # --- 하단 상태바 ---
        self.status_var = tk.StringVar(value="txt 파일을 선택해주세요.")
        ttk.Label(self.root, textvariable=self.status_var, relief="sunken", anchor="w").pack(
            fill="x", side="bottom", padx=10, pady=(0, 5)
        )

    def _open_file(self):
        path = filedialog.askopenfilename(
            title="카카오톡 채팅 로그 선택",
            filetypes=[("텍스트 파일", "*.txt"), ("모든 파일", "*.*")],
        )
        if path:
            self.input_file = path
            display = os.path.basename(path)
            self.file_label.config(text=display, foreground="black")
            self.status_var.set(f"파일 로드됨: {display}  |  '분석하기' 버튼을 눌러주세요.")

    def _analyze(self):
        if not self.input_file:
            messagebox.showwarning("경고", "먼저 txt 파일을 선택해주세요.")
            return

        days = self.days_var.get()
        try:
            self.user_counter = parse_chat_log(self.input_file, days=days)
        except Exception as e:
            messagebox.showerror("오류", f"파일 분석 중 오류 발생:\n{e}")
            return

        if not self.user_counter:
            messagebox.showinfo("결과 없음", f"최근 {days}일 내 메시지가 없습니다.")
            return

        # 테이블 갱신
        self.tree.delete(*self.tree.get_children())
        for rank, (user, count) in enumerate(self.user_counter.most_common(), 1):
            self.tree.insert("", "end", values=(rank, user, count))

        total_users = len(self.user_counter)
        total_msgs = sum(self.user_counter.values())
        self.tree.insert("", "end", values=("", "합계", total_msgs))

        self.status_var.set(f"분석 완료: {total_users}명, 총 {total_msgs}개 메시지 (최근 {days}일)")

    def _save_excel(self):
        if not self.user_counter:
            messagebox.showwarning("경고", "먼저 분석을 실행해주세요.")
            return

        path = filedialog.asksaveasfilename(
            title="엑셀 파일 저장",
            defaultextension=".xlsx",
            initialfile="채팅통계.xlsx",
            filetypes=[("Excel 파일", "*.xlsx")],
        )
        if not path:
            return

        try:
            create_excel(self.user_counter, path, days=self.days_var.get())
            messagebox.showinfo("완료", f"엑셀 파일 저장 완료!\n{path}")
            self.status_var.set(f"저장됨: {os.path.basename(path)}")
        except Exception as e:
            messagebox.showerror("오류", f"저장 중 오류 발생:\n{e}")


def main():
    root = tk.Tk()
    ChatStatsApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
